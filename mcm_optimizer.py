from pathlib import Path
from openpyxl import load_workbook


# ============================================================
# MATRIX CHAIN MULTIPLICATION FOR MEDICAL IMAGE OPTIMIZATION
# ============================================================

INPUT_FILE = "Medical_Image_Matrices.xlsx"
OUTPUT_FILE = "Medical_Image_Matrices_MCM.xlsx"


# ------------------------------------------------------------
# MCM Dynamic Programming
# ------------------------------------------------------------

def matrix_chain_order(p):
    """
    p represents matrix dimensions.

    For p = [32, 32, 16, 8, 4]:

    A1 = 32 x 32
    A2 = 32 x 16
    A3 = 16 x 8
    A4 = 8 x 4
    """

    n = len(p) - 1

    dp = [[0 for _ in range(n)] for _ in range(n)]
    split = [[0 for _ in range(n)] for _ in range(n)]

    for chain_length in range(2, n + 1):

        for i in range(n - chain_length + 1):

            j = i + chain_length - 1

            dp[i][j] = float("inf")

            for k in range(i, j):

                cost = (
                    dp[i][k]
                    + dp[k + 1][j]
                    + p[i] * p[k + 1] * p[j + 1]
                )

                if cost < dp[i][j]:
                    dp[i][j] = cost
                    split[i][j] = k

    return dp, split


# ------------------------------------------------------------
# Get optimal parenthesization
# ------------------------------------------------------------

def get_order(split, i, j):

    if i == j:
        return f"A{i + 1}"

    k = split[i][j]

    left = get_order(split, i, k)
    right = get_order(split, k + 1, j)

    return f"({left} × {right})"


# ------------------------------------------------------------
# Cost of multiplying from left to right
# ------------------------------------------------------------

def left_to_right_cost(p):

    # ((A1 A2) A3) A4

    cost1 = p[0] * p[1] * p[2]

    result_rows = p[0]
    result_cols = p[2]

    cost2 = result_rows * result_cols * p[3]

    result_cols = p[3]

    cost3 = result_rows * result_cols * p[4]

    return cost1 + cost2 + cost3


# ------------------------------------------------------------
# Load existing workbook
# ------------------------------------------------------------

if not Path(INPUT_FILE).exists():

    print("ERROR: Medical_Image_Matrices.xlsx not found.")

else:

    wb = load_workbook(INPUT_FILE)

    # Remove old MCM sheet if it exists
    if "MCM Results" in wb.sheetnames:
        del wb["MCM Results"]

    ws = wb.create_sheet("MCM Results")

    # --------------------------------------------------------
    # Explanation section
    # --------------------------------------------------------

    ws.append(["MEDICAL IMAGE OPTIMIZATION USING MATRIX CHAIN MULTIPLICATION"])
    ws.append([])

    ws.append(["Image processing matrix chain"])
    ws.append(["A1", "32 x 32"])
    ws.append(["A2", "32 x 16"])
    ws.append(["A3", "16 x 8"])
    ws.append(["A4", "8 x 4"])
    ws.append([])

    ws.append([
        "These matrices represent stages in an image-processing pipeline."
    ])

    ws.append([])

    # --------------------------------------------------------
    # Results headings
    # --------------------------------------------------------

    ws.append([
        "Record",
        "Image Type",
        "Image Number",
        "Matrix Chain",
        "Left-to-Right Cost",
        "Optimal Order",
        "Optimal MCM Cost",
        "Scalar Multiplications Saved"
    ])

    # Matrix dimensions
    dimensions = [32, 32, 16, 8, 4]

    # Run MCM
    dp, split = matrix_chain_order(dimensions)

    optimal_order = get_order(
        split,
        0,
        len(dimensions) - 2
    )

    optimal_cost = dp[0][len(dimensions) - 2]

    normal_cost = left_to_right_cost(dimensions)

    saved = normal_cost - optimal_cost

    # --------------------------------------------------------
    # Process 50 MRI + 50 X-Ray records
    # --------------------------------------------------------

    record = 1

    for image_type in ["MRI", "X-Ray"]:

        for image_number in range(1, 51):

            ws.append([
                record,
                image_type,
                image_number,
                "A1 × A2 × A3 × A4",
                normal_cost,
                optimal_order,
                optimal_cost,
                saved
            ])

            record += 1

    # --------------------------------------------------------
    # Save workbook
    # --------------------------------------------------------

    wb.save(OUTPUT_FILE)

    print()
    print("==============================================")
    print("MCM OPTIMIZATION COMPLETED")
    print("==============================================")
    print("Total records      :", record - 1)
    print("MRI records        : 50")
    print("X-Ray records      : 50")
    print("Matrix chain       : 32x32, 32x16, 16x8, 8x4")
    print("Normal cost        :", normal_cost)
    print("Optimal order      :", optimal_order)
    print("Optimal MCM cost   :", optimal_cost)
    print("Multiplications saved:", saved)
    print("Output             :", OUTPUT_FILE)
    print("==============================================")